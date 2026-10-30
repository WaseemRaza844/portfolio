#!/usr/bin/env python3
"""Generate editable Word files beside their Markdown sources in faculty applications."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
APPLICATIONS = HERE / "job-applications" / "faculty"


def setup(doc, spec, app_name):
    name = spec["output"]
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11), Inches(8.5)
    sec.top_margin = Inches(.52 if name.startswith(("01_", "05_")) else (.60 if name.startswith("03_") else .76))
    sec.bottom_margin = Inches(.49 if name.startswith(("01_", "05_")) else (.56 if name.startswith("03_") else .68))
    sec.left_margin = sec.right_margin = Inches(.73 if name.startswith(("01_", "05_")) else (.75 if name.startswith("03_") else .83))
    sec.header_distance, sec.footer_distance = Inches(.28), Inches(.36)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.35 if name.startswith("01_") else (9.15 if name.startswith("05_") else (9.5 if name.startswith("03_") else 10.25)))
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3.4 if name.startswith("01_") else (2.8 if name.startswith("05_") else (3.9 if name.startswith("03_") else 5.3)))
    normal.paragraph_format.line_spacing = 1.02 if name.startswith(("01_", "03_", "05_")) else 1.06
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name, title.font.size = "Aptos Display", Pt(16.2)
    title.font.bold, title.font.color.rgb = True, RGBColor(13, 67, 103)
    title.paragraph_format.space_after = Pt(5 if name.startswith(("01_", "05_")) else 8)
    title.paragraph_format.keep_with_next = True

    heading = styles["Heading 1"]
    heading.font.name, heading.font.size = "Aptos", Pt(11.2)
    heading.font.bold, heading.font.color.rgb = True, RGBColor(15, 80, 116)
    heading.paragraph_format.space_before = Pt(5 if name.startswith(("03_", "05_")) else 9)
    heading.paragraph_format.space_after = Pt(2 if name.startswith(("03_", "05_")) else 3)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.keep_together = True

    if "Heading 2" in styles:
        styles["Heading 2"].font.name = "Aptos"
        styles["Heading 2"].font.size = Pt(10)
        styles["Heading 2"].font.bold = True
        styles["Heading 2"].font.color.rgb = RGBColor(47, 78, 101)
        styles["Heading 2"].paragraph_format.keep_with_next = True

    note = styles.add_style("Application Note", WD_STYLE_TYPE.PARAGRAPH)
    note.base_style = normal
    note.font.size, note.font.color.rgb = Pt(8.2), RGBColor(91, 103, 119)
    note.paragraph_format.space_after = Pt(3 if name.startswith("05_") else 5)
    note.paragraph_format.keep_with_next = True

    doc.core_properties.title = spec.get("title", name.replace("_", " "))
    doc.core_properties.subject = f"Faculty application: {app_name}"
    doc.core_properties.author = "Waseem Raza"


def put_text(paragraph, text):
    """Render simple Markdown bold and italic spans as Word runs."""
    pattern = r"(\*\*.*?\*\*|\*[^*]+\*)"
    for chunk in re.split(pattern, text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


def fields_table(doc, fields):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width, table.columns[1].width = Inches(1.85), Inches(4.85)
    for label, value in fields:
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F5F8FA" if index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
        left, right = cells
        left_run = left.paragraphs[0].add_run(label)
        left_run.bold = True
        left_run.font.size = Pt(9)
        right_run = right.paragraphs[0].add_run(value)
        right_run.font.size = Pt(9)
        left.paragraphs[0].paragraph_format.space_after = Pt(1)
        right.paragraphs[0].paragraph_format.space_after = Pt(1)
    doc.add_paragraph("").paragraph_format.space_after = Pt(1)


def add_block(doc, block, spec):
    text = block.strip()
    if not text:
        return
    if text.startswith("# "):
        put_text(doc.add_paragraph(style="Title"), text[2:])
        return
    if text.startswith("## "):
        put_text(doc.add_paragraph(style="Heading 1"), text[3:])
        return
    if text.startswith("### "):
        put_text(doc.add_paragraph(style="Heading 2"), text[4:])
        return
    if text.startswith("> "):
        put_text(doc.add_paragraph(style="Application Note"), text[2:])
        return

    lines = text.splitlines()
    # The references sheet uses compact two-column editable tables.
    if spec.get("layout") == "references" and all(line.strip().startswith("- **") for line in lines):
        fields = []
        for line in lines:
            match = re.match(r"-\s+\*\*(.*?):\*\*\s*(.+)", line.strip())
            if match:
                fields.append(match.groups())
        if fields:
            fields_table(doc, fields)
            return

    if all(re.match(r"^\s*(?:- |\* |\d+\. |- \[[ xX]\] )", line) for line in lines):
        for line in lines:
            value = line.strip()
            if value.startswith("- [ ] "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(.1)
                p.paragraph_format.space_after = Pt(2.3)
                put_text(p, "☐ " + value[6:])
            elif value.startswith("- [x] ") or value.startswith("- [X] "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(.1)
                p.paragraph_format.space_after = Pt(2.3)
                put_text(p, "☒ " + value[6:])
            elif value.startswith(("- ", "* ")):
                put_text(doc.add_paragraph(style="List Bullet"), value[2:])
            else:
                match = re.match(r"^\d+\. (.*)$", value)
                put_text(doc.add_paragraph(style="List Number"), match.group(1) if match else value)
        return

    p = doc.add_paragraph()
    put_text(p, text)


def build_application(app_folder, manifest_path):
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    created = []
    for spec in manifest.get("documents", []):
        source = app_folder / spec["source"]
        output = app_folder / spec["output"]
        if not source.is_file():
            raise FileNotFoundError(f"Missing application source: {source}")
        doc = Document()
        setup(doc, spec, app_folder.name)
        blocks = re.split(r"\n\s*\n", source.read_text(encoding="utf-8").strip())
        for block in blocks:
            add_block(doc, block, spec)
        doc.save(output)
        created.append(output)
        print(f"Created {output.relative_to(HERE)} ({output.stat().st_size} bytes)")
    return created


def main():
    if not APPLICATIONS.exists():
        print("No faculty application folders found.")
        return
    manifests = sorted(APPLICATIONS.glob("*/application-documents.json"))
    if not manifests:
        print("No application document manifests found.")
        return
    for manifest in manifests:
        build_application(manifest.parent, manifest)


if __name__ == "__main__":
    main()
