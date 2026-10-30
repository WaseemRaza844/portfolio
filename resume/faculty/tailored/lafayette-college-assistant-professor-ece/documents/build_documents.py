#!/usr/bin/env python3
"""Build editable DOCX documents from browser-editable Markdown sources."""
from __future__ import annotations
import base64
import json
from pathlib import Path
import re
import zlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "source"
WORD = ROOT / "word"
SEED = ROOT / "source.seed.txt"
FILENAMES = (
    "01_Lafayette_Cover_Letter",
    "02_Lafayette_Teaching_Statement",
    "03_Lafayette_Research_Statement",
    "04_Lafayette_Professional_References_Template",
    "05_Lafayette_Application_Checklist",
)


def restore_seed():
    """Initialize Markdown once; later edits to Markdown are authoritative."""
    missing = [n for n in FILENAMES if not (SOURCE / (n + '.md')).exists()]
    if not missing:
        return
    if not SEED.exists():
        raise FileNotFoundError(f"Missing sources: {missing}")
    packed = base64.b64decode(SEED.read_text(encoding='utf-8').strip())
    data = json.loads(zlib.decompress(packed))
    SOURCE.mkdir(parents=True, exist_ok=True)
    for name in missing:
        key = name + '.md'
        if key not in data:
            raise KeyError(key)
        (SOURCE / key).write_text(data[key], encoding='utf-8')


def setup(doc, name):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11), Inches(8.5)
    sec.top_margin = Inches(.52 if name.startswith(('01_','05_')) else (.60 if name.startswith('03_') else .76))
    sec.bottom_margin = Inches(.49 if name.startswith(('01_','05_')) else (.56 if name.startswith('03_') else .68))
    sec.left_margin = sec.right_margin = Inches(.73 if name.startswith(('01_','05_')) else (.75 if name.startswith('03_') else .83))
    sec.header_distance, sec.footer_distance = Inches(.28), Inches(.36)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(9.35 if name.startswith('01_') else (9.15 if name.startswith('05_') else (9.5 if name.startswith('03_') else 10.25)))
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3.4 if name.startswith('01_') else (2.8 if name.startswith('05_') else (3.9 if name.startswith('03_') else 5.3)))
    normal.paragraph_format.line_spacing = 1.02 if name.startswith(('01_','03_','05_')) else 1.06
    normal.paragraph_format.widow_control = True
    title = styles['Title']
    title.font.name, title.font.size = 'Aptos Display', Pt(16.2)
    title.font.bold, title.font.color.rgb = True, RGBColor(13,67,103)
    title.paragraph_format.space_after, title.paragraph_format.keep_with_next = Pt(5 if name.startswith(('01_','05_')) else 8), True
    heading = styles['Heading 1']
    heading.font.name, heading.font.size = 'Aptos', Pt(11.2)
    heading.font.bold, heading.font.color.rgb = True, RGBColor(15,80,116)
    heading.paragraph_format.space_before = Pt(5 if name.startswith(('03_','05_')) else 9)
    heading.paragraph_format.space_after = Pt(2 if name.startswith(('03_','05_')) else 3)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.keep_together = True
    note = styles.add_style('Lafayette Note', WD_STYLE_TYPE.PARAGRAPH)
    note.base_style = normal
    note.font.size, note.font.color.rgb = Pt(8.2), RGBColor(91,103,119)
    note.paragraph_format.space_after = Pt(3 if name.startswith('05_') else 5)
    note.paragraph_format.keep_with_next = True
    doc.core_properties.title = name.replace('_',' ')
    doc.core_properties.subject = 'Lafayette College Assistant Professor ECE Application'
    doc.core_properties.author = 'Waseem Raza'


def put_text(p, text):
    for i, line in enumerate(text.split('  \n')):
        if i:
            p.add_run('\n')
        for chunk in re.split(r'(\*\*.*?\*\*)', line):
            if not chunk:
                continue
            if chunk.startswith('**') and chunk.endswith('**'):
                p.add_run(chunk[2:-2]).bold = True
            else:
                p.add_run(chunk)


def fields_table(doc, fields):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width, table.columns[1].width = Inches(1.85), Inches(4.85)
    for label, value in fields:
        cells = table.add_row().cells
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F5F8FA' if c == cells[0] else 'FFFFFF')
            c._tc.get_or_add_tcPr().append(shade)
        a,b = cells
        ra = a.paragraphs[0].add_run(label); ra.bold=True; ra.font.size=Pt(9)
        rb = b.paragraphs[0].add_run(value); rb.font.size=Pt(9)
        a.paragraphs[0].paragraph_format.space_after = Pt(1)
        b.paragraphs[0].paragraph_format.space_after = Pt(1)
    doc.add_paragraph('').paragraph_format.space_after = Pt(1)


def build(name):
    data = (SOURCE / (name+'.md')).read_text(encoding='utf-8')
    blocks = re.split(r'\n\s*\n', data.strip())
    doc = Document()
    setup(doc,name)
    i=0
    while i<len(blocks):
        text=blocks[i].strip()
        if text.startswith('# '):
            put_text(doc.add_paragraph(style='Title'),text[2:])
        elif text.startswith('## '):
            put_text(doc.add_paragraph(style='Heading 1'),text[3:])
            if name.startswith('04_'):
                fields=[]
                while i+1<len(blocks) and blocks[i+1].startswith('- **'):
                    i+=1
                    for line in blocks[i].split('\n'):
                        match=re.match(r'-\s+\*\*(.*?):\*\*\s*(.+)',line)
                        if match: fields.append(match.groups())
                if fields: fields_table(doc,fields)
        elif text.startswith('> '):
            put_text(doc.add_paragraph(style='Lafayette Note'),text[2:])
        elif text.startswith('- **'):
            fields=[]
            for line in text.split('\n'):
                match=re.match(r'-\s+\*\*(.*?):\*\*\s*(.+)',line)
                if match: fields.append(match.groups())
            if fields: fields_table(doc,fields)
        else:
            p=doc.add_paragraph()
            put_text(p,text)
            if text.startswith('☐'):
                p.paragraph_format.left_indent=Inches(.1)
                p.paragraph_format.space_after=Pt(2.3)
                for run in p.runs:run.font.size=Pt(9.35)
        i+=1
    WORD.mkdir(parents=True,exist_ok=True)
    out=WORD/(name+'.docx')
    doc.save(out)
    print(f"Created {out.name} ({out.stat().st_size} bytes)")


if __name__ == '__main__':
    restore_seed()
    for name in FILENAMES:
        build(name)
